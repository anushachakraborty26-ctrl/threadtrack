"""
src/build_excel_tool.py — Generate the ThreadTrack risk-scoring workbook.

This builds output/ThreadTrack_Risk_Workbook.xlsx, the Excel side of Phase 6.
The workbook lets a merchandiser or planner enter purchase orders and see, for
each one, its delay risk and return risk — with the reasons and a recommended
action — all computed automatically.

WHY A GENERATOR SCRIPT (rather than hand-building the spreadsheet)
    The workbook can be rebuilt any time with `python src/build_excel_tool.py`,
    and the scoring logic has a single source of truth — this file — which
    reproduces src/rule_scorer.py exactly: the same 15 factors, the same
    weights, the same Low / Medium / High band thresholds.

THE FOUR SHEETS
    Read me     — what the tool is and how to use it
    Order book  — enter orders here; scores, bands, reasons, actions auto-fill
    Summary     — a live roll-up of the order book (the operations view)
    Rubric      — the 15 scoring factors and their weights, fully auditable

Run:  python src/build_excel_tool.py
"""

import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_PATH = "output/ThreadTrack_Risk_Workbook.xlsx"

FIRST_ROW = 2     # first order row in the Order book sheet
LAST_ROW = 151    # last pre-built row — 150 orders can be entered

# --- colours ---
NAVY = "2E5B8A"
LIGHT = "F2F5F8"
GREEN = "C6EFCE"
AMBER = "FFEB9C"
RED = "FFC7CE"
WHITE = "FFFFFF"
GREY_LINE = "D9D9D9"

THIN = Side(style="thin", color=GREY_LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- Order book columns: (header, width, kind) ---
COLUMNS = [
    ("PO ID",              12, "input"),
    ("Vendor cluster",     15, "input"),
    ("New vendor?",        12, "input"),
    ("Vendor reliability", 17, "input"),
    ("Fabric",             10, "input"),
    ("Order qty",          10, "input"),
    ("Destination tier",   15, "input"),
    ("Payment mode",       13, "input"),
    ("Season",             11, "input"),
    ("Delay score",        11, "computed"),
    ("Return score",       12, "computed"),
    ("Delay risk",         11, "computed"),
    ("Return risk",        11, "computed"),
    ("Key risk drivers",   42, "computed"),
    ("Recommended action", 54, "computed"),
]

# --- drop-down options for the input columns (column letter -> list) ---
DROPDOWNS = {
    "B": '"Tirupur,Bengaluru,Ludhiana,Delhi NCR"',
    "C": '"Yes,No"',
    "D": '"0.85 or above,Below 0.85"',
    "E": '"knit,woven"',
    "G": '"Tier-1,Tier-2,Tier-3"',
    "H": '"COD,Prepaid"',
    "I": '"normal,festive,monsoon"',
}

# --- five example orders, low risk through high risk ---
EXAMPLES = [
    ["PO-EX01", "Bengaluru", "No",  "0.85 or above", "knit",  180, "Tier-1", "Prepaid", "normal"],
    ["PO-EX02", "Tirupur",   "No",  "0.85 or above", "knit",  220, "Tier-2", "COD",     "normal"],
    ["PO-EX03", "Delhi NCR", "No",  "Below 0.85",    "woven", 300, "Tier-2", "COD",     "normal"],
    ["PO-EX04", "Tirupur",   "Yes", "Below 0.85",    "knit",  400, "Tier-3", "COD",     "festive"],
    ["PO-EX05", "Ludhiana",  "No",  "Below 0.85",    "woven", 380, "Tier-3", "COD",     "monsoon"],
]

# --- the rubric, for the Rubric sheet: (factor, delay points, return points) ---
RUBRIC = [
    ("Baseline — every order starts here",        "10",  "15"),
    ("New vendor",                                "+22", "+10"),
    ("Vendor cluster: Tirupur",                   "+14", "+3"),
    ("Vendor cluster: Bengaluru",                 "+4",  "0"),
    ("Vendor reliability below 0.85",             "+12", "+6"),
    ("Fabric: woven",                             "+12", "+5"),
    ("Order quantity above 350",                  "+8",  "0"),
    ("Destination: Tier-3",                       "+10", "+14"),
    ("Destination: Tier-2",                       "+5",  "+7"),
    ("Payment: COD",                              "+2",  "+22"),
    ("Season: festive",                           "+12", "+10"),
    ("Season: monsoon",                           "+8",  "+3"),
    ("Interaction: monsoon + woven",              "+10", "+3"),
    ("Interaction: festive + COD",                "0",   "+14"),
    ("Interaction: new vendor + Tirupur",         "+10", "+2"),
    ("Predicted delay 60+ raises return risk",    "-",   "+15"),
    ("Predicted delay 40-59 raises return risk",  "-",   "+8"),
]


# ----------------------------------------------------------------------------
# Formulas — each reproduces the rule scorer exactly, for one row r
# ----------------------------------------------------------------------------

def delay_formula(r):
    return (f'=IF($B{r}="","",MIN(100,10'
            f'+IF($C{r}="Yes",22,0)'
            f'+IF($B{r}="Tirupur",14,0)'
            f'+IF($B{r}="Bengaluru",4,0)'
            f'+IF($D{r}="Below 0.85",12,0)'
            f'+IF($E{r}="woven",12,0)'
            f'+IF($F{r}>350,8,0)'
            f'+IF($G{r}="Tier-3",10,0)'
            f'+IF($G{r}="Tier-2",5,0)'
            f'+IF($H{r}="COD",2,0)'
            f'+IF($I{r}="festive",12,0)'
            f'+IF($I{r}="monsoon",8,0)'
            f'+IF(AND($I{r}="monsoon",$E{r}="woven"),10,0)'
            f'+IF(AND($C{r}="Yes",$B{r}="Tirupur"),10,0)))')


def return_formula(r):
    return (f'=IF($B{r}="","",MIN(100,15'
            f'+IF($C{r}="Yes",10,0)'
            f'+IF($B{r}="Tirupur",3,0)'
            f'+IF($D{r}="Below 0.85",6,0)'
            f'+IF($E{r}="woven",5,0)'
            f'+IF($G{r}="Tier-3",14,0)'
            f'+IF($G{r}="Tier-2",7,0)'
            f'+IF($H{r}="COD",22,0)'
            f'+IF($I{r}="festive",10,0)'
            f'+IF($I{r}="monsoon",3,0)'
            f'+IF(AND($I{r}="monsoon",$E{r}="woven"),3,0)'
            f'+IF(AND($I{r}="festive",$H{r}="COD"),14,0)'
            f'+IF(AND($C{r}="Yes",$B{r}="Tirupur"),2,0)'
            f'+IF($J{r}>=60,15,IF($J{r}>=40,8,0))))')


def band_formula(score_col, r):
    return (f'=IF(${score_col}{r}="","",'
            f'IF(${score_col}{r}<=35,"Low",'
            f'IF(${score_col}{r}<=65,"Medium","High")))')


def drivers_formula(r):
    joined = (f'_xlfn.TEXTJOIN(", ",TRUE,'
              f'IF($C{r}="Yes","New vendor",""),'
              f'IF($B{r}="Tirupur","Tirupur cluster",""),'
              f'IF($B{r}="Bengaluru","Bengaluru cluster",""),'
              f'IF($D{r}="Below 0.85","Low vendor reliability",""),'
              f'IF($E{r}="woven","Woven fabric",""),'
              f'IF($F{r}>350,"Large order (350+)",""),'
              f'IF($G{r}="Tier-3","Tier-3 destination",""),'
              f'IF($G{r}="Tier-2","Tier-2 destination",""),'
              f'IF($H{r}="COD","COD payment",""),'
              f'IF($I{r}="festive","Festive season",""),'
              f'IF($I{r}="monsoon","Monsoon season",""))')
    return f'=IF($B{r}="","",IF({joined}="","Baseline risk only",{joined}))'


def action_formula(r):
    return (f'=IF($B{r}="","",'
            f'IF(AND($L{r}="High",$M{r}="High"),'
            f'"Escalate — high delay AND return risk: alternate vendor, add buffer, check sizing and payment",'
            f'IF($L{r}="High",'
            f'"High delay risk — add a lead-time buffer or use a faster cluster",'
            f'IF($M{r}="High",'
            f'"High return risk — tighten sizing guidance; consider a prepaid incentive",'
            f'IF(OR($L{r}="Medium",$M{r}="Medium"),'
            f'"Moderate risk — keep this order on the watch list",'
            f'"Low risk — no action needed")))))')


# ----------------------------------------------------------------------------
# Sheet builders
# ----------------------------------------------------------------------------

def build_readme(ws):
    ws.column_dimensions["A"].width = 100
    lines = [
        ("ThreadTrack — Order Risk Workbook", 18, True, NAVY),
        ("", 11, False, None),
        ("WHAT THIS IS", 12, True, NAVY),
        ("A working tool for merchandisers and planners. Enter a purchase order and the "
         "workbook instantly estimates two risks — the chance it is delivered late, and "
         "the chance it is returned — each with the reasons behind it and a recommended "
         "action.", 11, False, None),
        ("", 11, False, None),
        ("HOW TO USE IT", 12, True, NAVY),
        ("1.  Open the 'Order book' sheet.", 11, False, None),
        ("2.  Fill one row per order. The shaded columns (Delay score onward) fill in by "
         "themselves — do not type in them.", 11, False, None),
        ("3.  Most fields are drop-downs: click the cell and pick a value.", 11, False, None),
        ("4.  Read the 'Recommended action' column for what to do about a risky order.", 11, False, None),
        ("5.  The 'Summary' sheet rolls the whole order book up automatically.", 11, False, None),
        ("", 11, False, None),
        ("WHAT THE COLOURS MEAN", 12, True, NAVY),
        ("Green = Low risk (score 0-35).   Amber = Medium risk (36-65).   "
         "Red = High risk (66-100).", 11, False, None),
        ("", 11, False, None),
        ("HOW THE SCORES ARE CALCULATED", 12, True, NAVY),
        ("The scores come from the ThreadTrack rule scorer — 15 supply chain risk "
         "factors, each adding points to a delay score and a return score. Every factor "
         "and its weight is listed, in full, on the 'Rubric' sheet. Nothing is hidden.", 11, False, None),
        ("", 11, False, None),
        ("Five example orders are pre-filled in the Order book so the tool is not empty "
         "on opening. Overwrite them with your own orders, or delete those rows.", 10, False, "888888"),
    ]
    for i, (text, size, bold, color) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(size=size, bold=bold, color=(color or "000000"))
        c.alignment = Alignment(wrap_text=True, vertical="top")


def build_order_book(ws):
    # header row
    for idx, (name, width, _kind) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        c = ws.cell(row=1, column=idx, value=name)
        c.font = Font(bold=True, color=WHITE, size=11)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 30

    computed_fill = PatternFill("solid", fgColor=LIGHT)

    # formula rows (every row gets its formulas; empty rows stay blank)
    for r in range(FIRST_ROW, LAST_ROW + 1):
        ws.cell(row=r, column=10, value=delay_formula(r))
        ws.cell(row=r, column=11, value=return_formula(r))
        ws.cell(row=r, column=12, value=band_formula("J", r))
        ws.cell(row=r, column=13, value=band_formula("K", r))
        ws.cell(row=r, column=14, value=drivers_formula(r))
        ws.cell(row=r, column=15, value=action_formula(r))
        for col in range(1, 16):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            if col >= 10:
                cell.fill = computed_fill
            if col in (12, 13):
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col in (14, 15):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")

    # five example orders
    for i, example in enumerate(EXAMPLES):
        r = FIRST_ROW + i
        for col, value in enumerate(example, start=1):
            ws.cell(row=r, column=col, value=value)

    # drop-downs on the input columns
    for col_letter, options in DROPDOWNS.items():
        dv = DataValidation(type="list", formula1=options, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}{FIRST_ROW}:{col_letter}{LAST_ROW}")

    # colour the two risk-band columns by value
    band_range = f"L{FIRST_ROW}:M{LAST_ROW}"
    ws.conditional_formatting.add(band_range, CellIsRule(
        operator="equal", formula=['"High"'], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add(band_range, CellIsRule(
        operator="equal", formula=['"Medium"'], fill=PatternFill("solid", fgColor=AMBER)))
    ws.conditional_formatting.add(band_range, CellIsRule(
        operator="equal", formula=['"Low"'], fill=PatternFill("solid", fgColor=GREEN)))

    ws.freeze_panes = "A2"


def build_summary(ws):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14

    ob = "'Order book'"
    rng = {c: f"{ob}!{c}{FIRST_ROW}:{c}{LAST_ROW}" for c in "BIJKLM"}

    def section(r, text):
        ws.cell(row=r, column=1, value=text).font = Font(size=12, bold=True, color=NAVY)

    def stat(r, label, formula):
        ws.cell(row=r, column=1, value=label).font = Font(size=11)
        ws.cell(row=r, column=2, value=formula).font = Font(size=11, bold=True)

    ws.cell(row=1, column=1, value="Order book — live summary").font = Font(
        size=14, bold=True, color=NAVY)

    section(3, "Overview")
    stat(4, "Orders entered", f"=COUNTA({rng['B']})")
    stat(5, "Average delay score", f'=IFERROR(ROUND(AVERAGE({rng["J"]}),1),"-")')
    stat(6, "Average return score", f'=IFERROR(ROUND(AVERAGE({rng["K"]}),1),"-")')
    stat(7, "High-risk orders (delay or return)",
         f'=SUMPRODUCT(--((({rng["L"]}="High")+({rng["M"]}="High"))>0))')

    section(9, "Delay risk breakdown")
    stat(10, "Low",    f'=COUNTIF({rng["L"]},"Low")')
    stat(11, "Medium", f'=COUNTIF({rng["L"]},"Medium")')
    stat(12, "High",   f'=COUNTIF({rng["L"]},"High")')

    section(14, "Return risk breakdown")
    stat(15, "Low",    f'=COUNTIF({rng["M"]},"Low")')
    stat(16, "Medium", f'=COUNTIF({rng["M"]},"Medium")')
    stat(17, "High",   f'=COUNTIF({rng["M"]},"High")')

    section(19, "Average delay score by cluster")
    for i, cluster in enumerate(["Tirupur", "Bengaluru", "Ludhiana", "Delhi NCR"]):
        stat(20 + i, cluster,
             f'=IFERROR(ROUND(AVERAGEIF({rng["B"]},"{cluster}",{rng["J"]}),1),"-")')

    section(25, "Average return score by season")
    for i, season in enumerate(["normal", "festive", "monsoon"]):
        stat(26 + i, season,
             f'=IFERROR(ROUND(AVERAGEIF({rng["I"]},"{season}",{rng["K"]}),1),"-")')


def build_rubric(ws):
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15

    ws.cell(row=1, column=1, value="The 15-factor scoring rubric").font = Font(
        size=14, bold=True, color=NAVY)
    intro = ws.cell(row=2, column=1, value=(
        "Every order starts at the baseline and gains points for each risk factor that "
        "applies. Both scores are capped at 100. This is the exact logic the Order book "
        "uses."))
    intro.font = Font(size=10, italic=True)
    intro.alignment = Alignment(wrap_text=True)

    for i, head in enumerate(["Risk factor", "Delay points", "Return points"], start=1):
        c = ws.cell(row=4, column=i, value=head)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER

    for i, (factor, delay_pts, return_pts) in enumerate(RUBRIC):
        r = 5 + i
        ws.cell(row=r, column=1, value=factor).font = Font(size=11)
        ws.cell(row=r, column=2, value=delay_pts).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=return_pts).alignment = Alignment(horizontal="center")
        shade = PatternFill("solid", fgColor=LIGHT) if i % 2 == 0 else None
        for col in range(1, 4):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            if shade:
                cell.fill = shade


def main():
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    build_readme(readme)
    build_order_book(wb.create_sheet("Order book"))
    build_summary(wb.create_sheet("Summary"))
    build_rubric(wb.create_sheet("Rubric"))
    wb.save(OUTPUT_PATH)

    print(f"Workbook written: {OUTPUT_PATH}")
    print("  Sheets: Read me, Order book, Summary, Rubric")
    print(f"  Order book: {LAST_ROW - FIRST_ROW + 1} scoreable rows, "
          f"{len(EXAMPLES)} example orders pre-filled")


if __name__ == "__main__":
    main()
